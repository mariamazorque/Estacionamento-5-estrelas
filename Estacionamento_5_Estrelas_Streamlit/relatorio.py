"""Planilha Excel do fechamento mensal."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from dados import placa_fmt, MODALIDADES, PAGAMENTOS, agora, br, dia_semana, mes, rotulo, data_valida

MOEDA = '"R$" #,##0.00'
AZUL = "073449"
DOURADO = "F6E7B4"


def _aba(ws, linhas, colunas_moeda, larguras, cabecalho=True):
    for linha in linhas:
        ws.append(linha)
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            if cell.column - 1 in colunas_moeda and isinstance(cell.value, (int, float)):
                cell.number_format = MOEDA
    if cabecalho:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=AZUL)
        ws.freeze_panes = "A2"


def gerar_excel(t: dict, comp: str) -> bytes:
    d = mes(t, comp)
    r = lambda c: round(c / 100, 2)
    wb = Workbook()

    # ---- Resumo
    ws = wb.active
    ws.title = "Resumo"
    linhas = [["Estacionamento 5 Estrelas"], [f"Fechamento de {rotulo(comp)}"], ["Gerado em", br(agora()[:10]) + " " + agora()[11:16]], [],
              ["RECEITAS", "Valor"], ["Avulsos em dinheiro", r(d["dinheiro"])], ["Avulsos em PIX", r(d["pix"])],
              ["Avulsos em cartão", r(d["cartao"])], ["Mensalistas", r(d["mensal"])], ["Total de receitas", r(d["receita"])], []]
    if d["tem_rec"]:
        linhas += [["POR MODALIDADE", "Valor"]] + [[m, r(d["mods"][m])] for m in MODALIDADES] \
            + [["Mensal", r(sum(p["valor"] for p in d["pays"]))], []]
    elif d["manual"]:
        linhas += [["Receitas avulsas informadas à mão no fechamento"], []]
    extra = {"Dinheiro": d["m_dinheiro"], "PIX": d["m_pix"], "Cartão": d["m_cartao"]}
    base = {"Dinheiro": d["dinheiro"], "PIX": d["pix"], "Cartão": d["cartao"]}
    linhas += [["RECEBIDO POR FORMA", "Valor"]] + [[f"{p} (avulsos + mensalidades)", r(base[p] + extra[p])] for p in PAGAMENTOS] + [[]]
    cats = sorted(d["cats"].items(), key=lambda x: -x[1])
    linhas += [["DESPESAS", "Valor"]] + ([[c, r(v)] for c, v in cats] or [["Nenhuma despesa lançada", 0]]) \
        + [["Total de despesas", r(d["despesas"])], [], ["RESULTADO DO MÊS", r(d["resultado"])]]
    _aba(ws, linhas, [1], [44, 18], cabecalho=False)
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"].font = Font(bold=True, size=12)
    for row in ws.iter_rows(min_row=4):
        a = row[0]
        if isinstance(a.value, str) and a.value.isupper():
            a.font = Font(bold=True, color=AZUL)
            if a.value == "RESULTADO DO MÊS":
                for c in row[:2]:
                    c.fill = PatternFill("solid", fgColor=DOURADO)
                    c.font = Font(bold=True)
        if a.value in ("Total de receitas", "Total de despesas"):
            for c in row[:2]:
                c.font = Font(bold=True)

    # ---- Receitas
    nomes = {m["id"]: m["nome"] for m in t["mensalistas"]}
    itens = [(e["data"], e["modalidade"], "", e["pagamento"], e["valor"]) for e in d["ent"]] + \
            [(p.get("data") or "", "Mensal", (p.get("nome") or nomes.get(p["mensalista_id"], "")) + (" (férias)" if p.get("ferias") else ""), p.get("pagamento") or "", p["valor"]) for p in d["pays"]]
    itens.sort(key=lambda x: x[0])
    rec = [[br(x[0]), dia_semana(x[0]) if data_valida(x[0]) else "", x[1], x[2], x[3], r(x[4])] for x in itens]
    total = r(sum(x[4] for x in itens))
    _aba(wb.create_sheet("Receitas"), [["Data", "Dia", "Modalidade", "Mensalista", "Pagamento", "Valor"], *rec, [], ["", "", "", "", "Total", total]],
         [5], [12, 10, 12, 26, 12, 14])

    # ---- Mensalistas
    pagos = {p["mensalista_id"]: p for p in d["pays"]}
    subs = [m for m in t["mensalistas"] if m["id"] in pagos or (m["ativo"] and (m.get("inicio") or "")[:7] <= comp)]
    linhas = []
    for m in sorted(subs, key=lambda m: m["nome"].lower()):
        p = pagos.get(m["id"])
        linhas.append([m["nome"], m.get("veiculo") or "", m.get("modelo") or "", placa_fmt(m.get("placa")), m.get("plano") or "", r(m["valor"]), ("Pago (férias)" if p.get("ferias") else "Pago") if p else "Pendente",
                       br(p.get("data")) if p else "", (p.get("pagamento") or "") if p else "", r(p["valor"]) if p else 0])
    ids = {m["id"] for m in t["mensalistas"]}
    for p in d["pays"]:
        if p["mensalista_id"] not in ids:
            linhas.append([p.get("nome") or "Mensalista (cadastro excluído)", "", "", "", "", "", "Pago (férias)" if p.get("ferias") else "Pago", br(p.get("data")), p.get("pagamento") or "", r(p["valor"])])
    _aba(wb.create_sheet("Mensalistas"), [["Mensalista", "Veículo", "Modelo", "Placa", "Plano", "Valor mensal", "Status", "Data pgto.", "Forma", "Valor pago"],
                                          *linhas, [], ["", "", "", "", "", "", "", "", "Total recebido", r(sum(p["valor"] for p in d["pays"]))]],
         [5, 9], [28, 10, 18, 11, 20, 14, 14, 12, 12, 14])

    # ---- Despesas
    ex = [[br(e["data"]), e["categoria"], e.get("descricao") or "", r(e["valor"])] for e in sorted(d["exp"], key=lambda e: e["data"])]
    _aba(wb.create_sheet("Despesas"), [["Data", "Categoria", "Descrição", "Valor"], *ex, [], ["", "", "Total", r(d["despesas"])]],
         [3], [12, 16, 32, 14])

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="center")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
